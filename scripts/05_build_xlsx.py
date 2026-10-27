"""data/listings_enriched.json → out/Paris_konaklama_v2.xlsx (skora göre sıralı tek sayfa + düşük yorumlar sayfası)

Kullanım: py scripts/05_build_xlsx.py
"""
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
BUDGET_TRY = 65000
HEAD = PatternFill("solid", fgColor="142033")
WARN = PatternFill("solid", fgColor="FBE3D6")

COLUMNS = [  # (başlık, genişlik, değer)
    ("Sıra", 6, lambda r: r["rank"]),
    ("Skor", 7, lambda r: r["score"]),
    ("ID", 6, lambda r: r["id"]),
    ("Platform", 9, lambda r: r["platform"]),
    ("İlan", 34, lambda r: r["title"]),
    ("Toplam TL (ödeme ekranı, vergi dahil, 9 gece)", 16, lambda r: r["price_total_try"]),
    ("İlan sayfasındaki TL", 12, lambda r: r.get("price_listed_try")),
    ("Yatak odası", 9, lambda r: r["bedrooms"]),
    ("2 ayrı oda", 12, lambda r: r["rooms_status"]),
    ("En fazla kişi", 8, lambda r: r["max_guests"]),
    ("Yatak düzeni", 36, lambda r: r.get("bed_layout")),
    ("Puan", 7, lambda r: r["rating"]),
    ("Ölçek", 6, lambda r: r["rating_scale"]),
    ("Yorum sayısı", 8, lambda r: r["review_count"]),
    ("3★ altı yorum (2 yıl)", 9, lambda r: None if r["low_reviews"] is None else len(r["low_reviews"])),
    ("Mines'e süre (dk, kapıdan kapıya)", 11, lambda r: r.get("transit_min")),
    ("Mines'e km (kuş uçuşu)", 10, lambda r: r.get("dist_mines_km")),
    ("En yakın istasyon", 22, lambda r: (r.get("nearest_stations") or [{}])[0].get("name")),
    ("Hatlar", 14, lambda r: ", ".join((r.get("nearest_stations") or [{}])[0].get("lines") or [])),
    ("İstasyona m (kuş uçuşu)", 10, lambda r: (r.get("nearest_stations") or [{}])[0].get("dist_m")),
    ("Konum", 10, lambda r: "kesin" if r["location_precision"] == "exact" else "yaklaşık"),
    ("Adres / yer", 30, lambda r: r.get("address")),
    ("İptal", 40, lambda r: r.get("cancellation")),
    ("Giriş saati", 14, lambda r: r.get("checkin_time")),
    ("Not", 45, lambda r: r.get("notes")),
    ("İlan", 12, lambda r: r["url"]),
    ("Mines'e rota", 12, lambda r: r.get("gmaps_transit_url")),
]


def main():
    data = json.loads((ROOT / "data" / "listings_enriched.json").read_text(encoding="utf-8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "İlanlar"
    ws.append([c[0] for c in COLUMNS])
    for i, (_, w, _) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        cell = ws.cell(1, i)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 42

    link_cols = {i for i, c in enumerate(COLUMNS, 1) if c[0] in ("İlan", "Mines'e rota") and i > 5}
    for r in data:
        ws.append([c[2](r) for c in COLUMNS])
        row = ws.max_row
        for i in link_cols:
            cell = ws.cell(row, i)
            if cell.value:
                cell.hyperlink, cell.value, cell.style = cell.value, "aç ↗", "Hyperlink"
        for i in range(1, len(COLUMNS) + 1):
            ws.cell(row, i).alignment = Alignment(vertical="top", wrap_text=COLUMNS[i - 1][1] >= 22)
        ws.cell(row, 6).number_format = ws.cell(row, 7).number_format = '#,##0" ₺"'
        if r["over_budget"]:
            ws.cell(row, 6).fill = WARN
        if r["rooms_status"] != "onaylı":
            ws.cell(row, 9).fill = WARN

    n = ws.max_row
    ws.conditional_formatting.add(f"B2:B{n}", ColorScaleRule(start_type="num", start_value=20, start_color="F2B8B5",
                                                             mid_type="num", mid_value=50, mid_color="FFE9A8",
                                                             end_type="num", end_value=80, end_color="B7E1C1"))
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{n}"

    rv = wb.create_sheet("Düşük puanlı yorumlar")
    rv.append(["ID", "İlan", "Tarih", "Yıldız", "Yorum (İngilizce çeviri)"])
    for i, w in enumerate((6, 34, 11, 7, 110), 1):
        rv.column_dimensions[get_column_letter(i)].width = w
        rv.cell(1, i).font, rv.cell(1, i).fill = Font(bold=True, color="FFFFFF"), HEAD
    for r in data:
        for x in r["low_reviews"] or []:
            rv.append([r["id"], r["title"], x["date"], x["stars"], x["text"]])
            rv.cell(rv.max_row, 5).alignment = Alignment(wrap_text=True, vertical="top")
    rv.freeze_panes = "A2"

    info = wb.create_sheet("Açıklama")
    for line in [
        f"Oluşturulma: {date.today().isoformat()} · 13–22 Kasım 2026, 9 gece, 3 yetişkin · bütçe {BUDGET_TRY:,} TL".replace(",", "."),
        "Toplam TL: ödeme ekranında görülen vergi dahil fiyat (kullanıcı, 27.10.2026; data/price_checked.csv). Airbnb'de ödeme ekranı ilan sayfasından çoğunlukla ~%15 yüksek.",
        "Skor (0–100), önem sırasıyla: 2 ayrı yatak odası 40 (salondaki çekyat sayılmaz; oda dökümü yoksa 32, şüpheliyse 18), puan 15 + yorumlar 10 (her 3★ altı yorum −2,5, izinsiz giriş iddiası −5), Mines'e süre 20 (15 dk tam, 55 dk sıfır, 1 saat üstü −5), fiyat 15. Fiyatı olmayan ilan en çok 45.",
        "Airbnb konumu platformun kaydırdığı yaklaşık konum (~300 m); Booking adresi kesin. Mesafeler kuş uçuşu.",
        "Boş hücre = sayfadan okunamadı. Ödeme ekranında fiyatı ve iptal tarihini mutlaka teyit et.",
        "Not sütunu orijinal tablodan (Paris_27_Ilan.xlsx) aktarıldı.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 150

    out = ROOT / "out" / "Paris_konaklama_v2.xlsx"
    wb.save(out)
    print(f"{out} — {len(data)} ilan")


if __name__ == "__main__":
    main()
