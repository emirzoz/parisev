"""Orijinal GPT tablosundan ilan URL'lerini çıkarıp data/listings_input.csv yazar.

Kullanım: py scripts/01_extract_input.py
Orijinal xlsx'e dokunmaz.
"""
import csv
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "Paris_27_Ilan.xlsx"
OUT = ROOT / "data" / "listings_input.csv"

HEADER_ROW = 6  # "Platform | İlan | ..." satırı
COLS = {"platform": 1, "name": 2, "base": 3, "tax": 4, "rooms": 6, "rating": 7, "notes": 8, "url": 9}


def listing_key(url: str) -> str:
    m = re.search(r"/rooms/(\d+)", url)
    if m:
        return m.group(1)
    m = re.search(r"/hotel/fr/([^./?]+)", url)
    return m.group(1) if m else ""


def main() -> None:
    ws = openpyxl.load_workbook(SRC)["Gonderdigin ilanlar"]
    rows = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        url = ws.cell(r, COLS["url"]).value
        if not url:
            continue
        rows.append(
            {
                "id": f"L{len(rows) + 1:02d}",
                "platform": ws.cell(r, COLS["platform"]).value,
                "url": url.strip(),
                "listing_key": listing_key(url),
                "gpt_name": ws.cell(r, COLS["name"]).value,
                "gpt_price_base_try": ws.cell(r, COLS["base"]).value,
                "gpt_tax_try": ws.cell(r, COLS["tax"]).value,
                "gpt_rooms": ws.cell(r, COLS["rooms"]).value,
                "gpt_rating": ws.cell(r, COLS["rating"]).value,
                "notes": ws.cell(r, COLS["notes"]).value,
            }
        )
    OUT.parent.mkdir(exist_ok=True)
    with OUT.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    by_platform = {}
    for row in rows:
        by_platform[row["platform"]] = by_platform.get(row["platform"], 0) + 1
    print(f"{len(rows)} ilan -> {OUT}  {by_platform}")


if __name__ == "__main__":
    main()
